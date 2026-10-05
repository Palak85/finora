import csv
from io import BytesIO
from decimal import Decimal
from datetime import datetime
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import permissions, status

from apps.accounts.models import Account
from apps.transactions.models import Transaction
from apps.budgets.models import Budget
from apps.investments.models import Investment
from apps.analytics.utils import calculate_total_balance, calculate_monthly_income, calculate_monthly_expenses

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ReportSummaryView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        user = request.user
        txns = Transaction.objects.filter(user=user)
        accounts = Account.objects.filter(user=user)
        budgets = Budget.objects.filter(user=user)
        investments = Investment.objects.filter(user=user)

        return Response({
            'success': True,
            'data': {
                'total_transactions': txns.count(),
                'total_accounts': accounts.count(),
                'total_budgets': budgets.count(),
                'total_investments': investments.count(),
                'available_reports': [
                    {'id': 'monthly_financial', 'name': 'Monthly Financial Report', 'formats': ['pdf', 'csv', 'excel']},
                    {'id': 'expense', 'name': 'Expense Report', 'formats': ['pdf', 'csv', 'excel']},
                    {'id': 'budget', 'name': 'Budget Adherence Report', 'formats': ['pdf', 'csv', 'excel']},
                    {'id': 'investment', 'name': 'Investment Portfolio Report', 'formats': ['pdf', 'csv', 'excel']},
                    {'id': 'net_worth', 'name': 'Net Worth Report', 'formats': ['pdf', 'csv', 'excel']},
                ]
            }
        })

class ExportReportCSVView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        report_type = request.query_params.get('type', 'monthly_financial')
        user = request.user
        
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="finora_{report_type}_report.csv"'

        writer = csv.writer(response)

        if report_type == 'expense':
            writer.writerow(['Date', 'Reference ID', 'Category', 'Payment Method', 'Description', 'Amount (INR)'])
            txns = Transaction.objects.filter(user=user, transaction_type='EXPENSE').order_by('-transaction_date')
            for t in txns:
                writer.writerow([t.transaction_date, t.reference_id, t.category, t.payment_method, t.description, t.amount])

        elif report_type == 'investment':
            writer.writerow(['Asset Name', 'Type', 'Quantity', 'Avg Buy Price', 'Invested Amount', 'Current Value', 'P&L', 'P&L %'])
            invs = Investment.objects.filter(user=user)
            for i in invs:
                writer.writerow([i.asset_name, i.asset_type, i.quantity, i.average_buy_price, i.invested_amount, i.current_value, i.profit_loss, i.profit_loss_percentage])

        elif report_type == 'budget':
            writer.writerow(['Category', 'Month/Year', 'Monthly Limit', 'Spent Amount', 'Status'])
            budgets = Budget.objects.filter(user=user)
            for b in budgets:
                writer.writerow([b.category, f"{b.month}/{b.year}", b.monthly_limit, b.spent_amount, 'SAFE'])

        else: # monthly_financial or net_worth
            writer.writerow(['Date', 'Reference ID', 'Type', 'Category', 'Description', 'Amount (INR)'])
            txns = Transaction.objects.filter(user=user).order_by('-transaction_date')
            for t in txns:
                writer.writerow([t.transaction_date, t.reference_id, t.transaction_type, t.category, t.description, t.amount])

        return response

class ExportReportExcelView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        report_type = request.query_params.get('type', 'monthly_financial')
        user = request.user

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Finora Report"

        header_font = Font(name='Arial', size=11, bold=True, color='F5F1E6')
        header_fill = PatternFill(start_color='0B0B0A', end_color='0B0B0A', fill_type='solid')
        gold_border = Side(style='thin', color='D4AF37')
        border_all = Border(left=gold_border, right=gold_border, top=gold_border, bottom=gold_border)

        if report_type == 'expense':
            headers = ['Date', 'Reference ID', 'Category', 'Payment Method', 'Description', 'Amount (INR)']
            ws.append(headers)
            txns = Transaction.objects.filter(user=user, transaction_type='EXPENSE').order_by('-transaction_date')
            for t in txns:
                ws.append([str(t.transaction_date), t.reference_id, t.category, t.payment_method, t.description, float(t.amount)])
        
        elif report_type == 'investment':
            headers = ['Asset Name', 'Type', 'Quantity', 'Avg Buy Price', 'Invested Amount', 'Current Value', 'P&L', 'P&L %']
            ws.append(headers)
            invs = Investment.objects.filter(user=user)
            for i in invs:
                ws.append([i.asset_name, i.asset_type, float(i.quantity), float(i.average_buy_price), float(i.invested_amount), float(i.current_value), float(i.profit_loss), float(i.profit_loss_percentage)])
        
        else:
            headers = ['Date', 'Reference ID', 'Type', 'Category', 'Description', 'Amount (INR)']
            ws.append(headers)
            txns = Transaction.objects.filter(user=user).order_by('-transaction_date')
            for t in txns:
                ws.append([str(t.transaction_date), t.reference_id, t.transaction_type, t.category, t.description, float(t.amount)])

        # Apply header styling
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="finora_{report_type}_report.xlsx"'
        return response

class ExportReportPDFView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        report_type = request.query_params.get('type', 'monthly_financial')
        user = request.user
        
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            textColor=colors.HexColor('#D4AF37'),
            spaceAfter=10
        )
        subtitle_style = ParagraphStyle(
            'SubTitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=11,
            textColor=colors.HexColor('#555555'),
            spaceAfter=20
        )

        elements.append(Paragraph("FINORA BANKING & PERSONAL FINANCE", title_style))
        elements.append(Paragraph(f"Official Statement: {report_type.replace('_', ' ').title()} Report | Generated for {user.full_name} ({user.email})", subtitle_style))
        elements.append(Spacer(1, 10))

        txns = Transaction.objects.filter(user=user).order_by('-transaction_date')[:30]
        data = [["Date", "Type", "Category", "Description", "Amount (INR)"]]

        for t in txns:
            data.append([
                str(t.transaction_date),
                t.transaction_type,
                t.category,
                t.description[:25],
                f"₹{t.amount}"
            ])

        t_table = Table(data, colWidths=[80, 80, 100, 180, 90])
        t_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0B0B0A')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#D4AF37')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('BACKGROUND', (0,1), (-1,-1), colors.HexColor('#F8F9FA')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ]))

        elements.append(t_table)
        doc.build(elements)

        pdf = buffer.getvalue()
        buffer.close()

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="finora_{report_type}_report.pdf"'
        response.write(pdf)
        return response
